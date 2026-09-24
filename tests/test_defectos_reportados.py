# © 2026 Martín Viera. Todos los derechos reservados.
"""
Plania · Tres defectos reportados por un usuario, fijados con tests
===================================================================
1. Sin python-docx la app se caía en cada pantalla con exportes
   (`ModuleNotFoundError: No module named 'docx'`), porque los bytes del Word
   se armaban en cada render aunque nadie apretara el botón.
2. Un Excel con el diccionario de datos en la primera hoja se rechazaba con
   «No pude mapear columnas obligatorias de productos: ['sku', 'precio']»
   aunque los productos estuvieran en otra hoja del mismo libro.
3. Los filtros de la barra lateral sólo aplicaban en algunas pestañas.

Todos los datos de estos tests son sintéticos, armados acá mismo.
"""
from __future__ import annotations

import ast
import os
import sys

import pandas as pd
import pytest

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
APP = os.path.join(RAIZ, "app", "app.py")


# ---------------------------------------------------------------------------
# 1. python-docx ausente
# ---------------------------------------------------------------------------
def _sin_docx(monkeypatch):
    # `None` en sys.modules hace que `import docx` levante ModuleNotFoundError,
    # que es exactamente lo que pasa en una PC sin la librería.
    for nombre in [m for m in sys.modules if m == "docx" or m.startswith("docx.")]:
        monkeypatch.delitem(sys.modules, nombre)
    monkeypatch.setitem(sys.modules, "docx", None)


def test_word_disponible_detecta_que_falta_python_docx(monkeypatch):
    from plania import exportes
    _sin_docx(monkeypatch)
    assert exportes.word_disponible() is False


def test_a_word_sin_python_docx_da_un_error_que_dice_como_arreglarlo(monkeypatch):
    from plania import exportes
    _sin_docx(monkeypatch)
    with pytest.raises(RuntimeError, match="pip install python-docx"):
        exportes.a_word("x", [("s", "t", pd.DataFrame({"a": [1]}))])


def test_pdf_y_excel_andan_sin_python_docx(monkeypatch):
    from plania import exportes
    _sin_docx(monkeypatch)
    secciones = [("Sección", "texto", pd.DataFrame({"sku": ["A1"], "precio": [10.0]}))]
    assert exportes.a_pdf("Informe", secciones)[:4] == b"%PDF"
    assert exportes.a_excel(secciones)[:2] == b"PK"


def test_la_app_no_arma_el_word_en_cada_render():
    """`a_word` no puede llamarse directo como argumento del botón: eso lo
    ejecuta en cada render. Tiene que ir diferido y detrás del chequeo."""
    fuente = open(APP, encoding="utf-8").read()
    assert "exportes.word_disponible()" in fuente
    assert "_diferido(lambda: exportes.a_word(" in fuente
    assert "download_button(t(\"comun.word\"), exportes.a_word(" not in fuente


def test_la_app_no_se_cae_sin_python_docx(monkeypatch):
    """La prueba de verdad: la app corre, sin python-docx, en las pantallas
    que tienen botones de exportar — y el de Word queda apagado."""
    from streamlit.testing.v1 import AppTest

    from plania import i18n, licencia

    # AppTest corre el script en ESTE proceso: la EULA y la licencia se fijan
    # acá para no depender de la config de la máquina que corre los tests.
    monkeypatch.setattr(licencia, "eula_aceptada", lambda: True)
    monkeypatch.setattr(licencia, "estado", lambda: {
        "modo": "demo", "dias_restantes": 7, "horas_restantes": 100})
    monkeypatch.setattr(licencia, "tiene", lambda _funcion: True)
    monkeypatch.syspath_prepend(RAIZ)
    _sin_docx(monkeypatch)

    at = AppTest.from_file(APP, default_timeout=180)
    at.run()
    assert not at.exception, [e.value for e in at.exception]
    for pagina in ("stock", "precios", "ofertas"):
        at.sidebar.radio[0].set_value(i18n.t(f"menu.{pagina}")).run()
        assert not at.exception, (pagina, [e.value for e in at.exception])
        apagados = [b.label for b in at.button if b.disabled]
        assert i18n.t("comun.word") in apagados, pagina


# ---------------------------------------------------------------------------
# 2. Excel de varias hojas con el diccionario adelante
# ---------------------------------------------------------------------------
def _libro(tmp_path, hojas: dict) -> str:
    ruta = tmp_path / "bases_sinteticas.xlsx"
    with pd.ExcelWriter(ruta, engine="openpyxl") as w:
        for nombre, df in hojas.items():
            df.to_excel(w, sheet_name=nombre, index=False)
    return str(ruta)


DICCIONARIO = pd.DataFrame({
    "Tabla": ["Productos", "Productos", "Ventas"],
    "Campo": ["Codigo", "Precio", "Fecha"],
    "Tipo": ["texto", "número", "fecha"],
    "Descripción": ["Código interno", "Precio de lista", "Fecha de emisión"],
})
PRODUCTOS = pd.DataFrame({
    "Codigo": ["A1", "A2", "A3"],
    "Descripcion": ["Yerba 1kg", "Azúcar 1kg", "Aceite 900ml"],
    "Precio": [210.0, 55.5, 120.0],
})


def test_elige_la_hoja_de_productos_y_no_el_diccionario(tmp_path):
    from plania import conectores
    ruta = _libro(tmp_path, {"Diccionario": DICCIONARIO, "Productos": PRODUCTOS})

    evaluacion = conectores.evaluar_hojas(ruta, "productos")
    assert conectores.mejor_hoja(evaluacion) == "Productos"
    assert evaluacion[0]["diccionario"] is True

    df = conectores.normalizar(conectores.leer_archivo(ruta, entidad="productos"),
                               "productos")
    assert list(df["sku"]) == ["A1", "A2", "A3"]
    assert df["precio"].tolist() == [210.0, 55.5, 120.0]


def test_la_hoja_se_puede_forzar(tmp_path):
    """Lo que usa el selector de la pantalla: la hoja elegida a mano gana."""
    from plania import conectores
    ruta = _libro(tmp_path, {"Diccionario": DICCIONARIO, "Productos": PRODUCTOS})
    df = conectores.leer_archivo(ruta, entidad="productos", hoja="Diccionario")
    assert "Campo" in df.columns


def test_si_ninguna_hoja_sirve_el_error_nombra_cada_hoja_y_lo_que_falta(tmp_path):
    from plania import conectores
    sin_precio = PRODUCTOS.drop(columns=["Precio"])
    ruta = _libro(tmp_path, {"Diccionario": DICCIONARIO, "Articulos": sin_precio})

    evaluacion = conectores.evaluar_hojas(ruta, "productos")
    texto = conectores.explicar_hojas(evaluacion, "productos")
    assert "Diccionario" in texto and "diccionario de datos" in texto
    assert "Articulos" in texto and "precio" in texto
    # Y no se inventa: la hoja elegida sigue fallando por el precio.
    with pytest.raises(ValueError, match="precio"):
        conectores.normalizar(conectores.leer_archivo(ruta, entidad="productos"),
                              "productos")


def test_precio_derivado_solo_a_pedido_y_anotado(tmp_path):
    from plania import conectores
    mercado = pd.DataFrame({
        "ProductoID": ["P1", "P1", "P2"],
        "Nombre": ["Uno", "Uno", "Dos"],
        "Unidades": [10, 30, 0],
        "VentasUSD": [100.0, 500.0, 50.0],
    })
    montos, unidades = conectores.candidatas_precio_derivado(mercado)
    assert montos == ["VentasUSD"] and unidades == ["Unidades"]

    # Sin pedirlo, no hay precio: el error sigue.
    with pytest.raises(ValueError, match="precio"):
        conectores.normalizar(mercado, "productos")

    d, nota = conectores.derivar_precio(mercado, "VentasUSD", "Unidades",
                                        col_clave="ProductoID")
    assert "DERIVADO" in nota and "VentasUSD" in nota and "Unidades" in nota
    precios = dict(zip(d["ProductoID"], d["precio"]))
    assert precios["P1"] == 15.0          # (100 + 500) / (10 + 30)
    assert pd.isna(precios["P2"])         # unidades en cero: vacío, no infinito
    assert len(d) == 2                    # una fila por producto


def test_ventas_toma_la_hoja_con_fecha_producto_y_unidades(tmp_path):
    from plania import conectores
    ventas = pd.DataFrame({"Fecha": ["2026-01-01", "2026-01-02"],
                           "ProductoID": ["A1", "A2"], "Unidades": [3, 4]})
    ruta = _libro(tmp_path, {"Diccionario": DICCIONARIO, "Productos": PRODUCTOS,
                             "Mercado": ventas})
    assert conectores.mejor_hoja(conectores.evaluar_hojas(ruta, "ventas")) == "Mercado"
    df = conectores.normalizar(conectores.leer_archivo(ruta, entidad="ventas"), "ventas")
    assert list(df["sku"]) == ["A1", "A2"]


def test_un_csv_no_pasa_por_la_seleccion_de_hojas(tmp_path):
    from plania import conectores
    ruta = tmp_path / "p.csv"
    PRODUCTOS.to_csv(ruta, index=False)
    assert conectores.hojas_de(str(ruta)) == []
    assert len(conectores.leer_archivo(str(ruta), entidad="productos")) == 3


# ---------------------------------------------------------------------------
# 3. Filtros en todas las pestañas
# ---------------------------------------------------------------------------
def _constante(nombre: str):
    for nodo in ast.parse(open(APP, encoding="utf-8").read()).body:
        if isinstance(nodo, ast.Assign) and any(
                isinstance(o, ast.Name) and o.id == nombre for o in nodo.targets):
            return nodo.value
    raise AssertionError(f"{nombre} no está en app/app.py")


def test_todas_las_pestanas_con_datos_filtran():
    menu = ast.literal_eval(_constante("MENU_CLAVES"))
    sin_datos = ast.literal_eval(_constante("PAGINAS_SIN_DATOS"))
    # PAGINAS_CON_FILTRO se deriva del menú: si alguien vuelve a escribirla
    # a mano, este test lo nota.
    con_filtro = eval(compile(ast.Expression(_constante("PAGINAS_CON_FILTRO")),  # noqa: S307
                              "app.py", "eval"),
                      {"MENU_CLAVES": menu, "PAGINAS_SIN_DATOS": sin_datos})
    assert set(sin_datos) == {"planes", "configuracion", "ayuda", "conectar_erp"}
    assert set(con_filtro) == set(menu) - set(sin_datos)
    for pagina in ("inicio", "panel_ejecutivo", "stock", "precios", "zonas",
                   "rutas", "ofertas", "copiloto"):
        assert pagina in con_filtro, pagina


def test_cada_pestana_filtrada_avisa_el_filtro():
    """Filtrar sin decirlo lleva a leer un recorte como si fuera el negocio."""
    fuente = open(APP, encoding="utf-8").read()
    bloques = fuente.split('elif pagina == "')
    inicio = fuente.split('if pagina == "inicio":', 1)[1].split("elif pagina ==", 1)[0]
    por_pagina = {"inicio": inicio}
    for b in bloques[1:]:
        clave, cuerpo = b.split('"', 1)
        por_pagina[clave] = cuerpo
    for pagina in ("inicio", "panel_ejecutivo", "stock", "precios", "zonas",
                   "rutas", "ofertas", "copiloto"):
        assert "_aviso_filtros()" in por_pagina[pagina], pagina


# ---------------------------------------------------------------------------
# «ID» por nombre de hoja + precio sacado de las ventas (Bases y diccionario)
# ---------------------------------------------------------------------------
from plania import conectores  # noqa: E402


def _libro_farma(ruta):
    import openpyxl
    wb = openpyxl.Workbook()
    d = wb.active
    d.title = "Diccionario"
    d.append(["Tabla", "Campo", "Tipo", "Descripción"])
    d.append(["Producto", "ID", "texto", "código"])
    p = wb.create_sheet("Producto")
    p.append(["ID", "Producto", "Molecula"])
    p.append(["P1", "Uno", "x"])
    p.append(["P2", "Dos", "y"])
    m = wb.create_sheet("Medico")
    m.append(["ID", "NombreMedico", "Ciudad"])
    m.append(["M1", "Médico 1", "Colonia"])
    v = wb.create_sheet("Visitas")
    v.append(["ID", "Fecha", "MedicoID", "ProductoID"])
    v.append([1, "2025-01-01", "M1", "P1"])
    mer = wb.create_sheet("Mercado")
    mer.append(["Fecha", "ProductoID", "Unidades", "VentasUSD"])
    mer.append(["2025-01-01", "P1", 10, 100.0])
    mer.append(["2025-02-01", "P1", 10, 300.0])
    mer.append(["2025-01-01", "P2", 0, 50.0])
    wb.save(ruta)


def test_un_ID_suelto_es_la_clave_si_la_hoja_se_llama_como_la_entidad(tmp_path):
    ruta = tmp_path / "libro.xlsx"
    _libro_farma(ruta)
    p = conectores.leer_archivo(str(ruta), "productos")
    assert p.attrs["hoja"] == "Producto"
    assert conectores.autodetectar_mapeo(p, "productos")["ID"] == "sku"
    c = conectores.leer_archivo(str(ruta), "clientes")
    mc = conectores.autodetectar_mapeo(c, "clientes")
    assert c.attrs["hoja"] == "Medico" and mc["ID"] == "cliente_id"
    assert mc["NombreMedico"] == "nombre"


def test_en_otra_hoja_un_ID_no_se_toma_como_clave():
    import pandas as pd
    df = pd.DataFrame({"ID": [1], "Fecha": ["2025-01-01"]})
    df.attrs["hoja"] = "Visitas"
    assert "ID" not in conectores.autodetectar_mapeo(df, "productos")
    df.attrs = {}
    assert "ID" not in conectores.autodetectar_mapeo(df, "clientes")


def test_precio_desde_ventas_es_el_realizado_y_viene_rotulado(tmp_path):
    ruta = tmp_path / "libro.xlsx"
    _libro_farma(ruta)
    v = conectores.leer_archivo(str(ruta), "ventas")
    precio, nota = conectores.precio_desde_ventas(v)
    assert precio["P1"] == 20.0                 # (100+300)/(10+10)
    assert "P2" not in precio.index             # unidades 0: no se inventa
    assert "DERIVADO" in nota
    p = conectores.leer_archivo(str(ruta), "productos")
    m = conectores.autodetectar_mapeo(p, "productos")
    out = conectores.normalizar(conectores.completar_precio(p, m, precio), "productos", m)
    assert out.set_index("sku").loc["P1", "precio"] == 20.0


def test_sin_monto_en_las_ventas_no_hay_precio_derivado():
    import pandas as pd
    v = pd.DataFrame({"fecha": ["2025-01-01"], "sku": ["P1"], "cantidad": [3]})
    assert conectores.precio_desde_ventas(v) == (None, "")


# ---------------------------------------------------------------------------
# «Hoja de ventas» en la hoja equivocada (Bases y diccionario, segundo reporte)
# ---------------------------------------------------------------------------
# El error decía a la vez «No pude mapear ventas: ['fecha', 'sku',
# 'cantidad']. Columnas disponibles: ['ID', 'Producto', …]» y «Mercado: tiene
# todas las obligatorias»: las ventas se leían de la hoja «Producto». De
# rebote, los productos caían por «falta precio», porque el precio se sacaba
# de esa misma hoja de ventas. Libro SINTÉTICO con la misma forma que el
# real: ocho hojas, mismos encabezados, valores inventados.
def _libro_ocho_hojas(ruta, con_mercado=True):
    import openpyxl
    wb = openpyxl.Workbook()
    d = wb.active
    d.title = "Diccionario"
    d.append(["Tabla", "Campo", "Tipo", "Descripción"])
    for campo in ("ID", "Producto", "Unidades", "VentasUSD"):
        d.append(["Mercado", campo, "texto", "sintético"])
    med = wb.create_sheet("Medico")
    med.append(["ID", "NombreMedico", "Especialidad", "Ciudad", "Segmento",
                "Potencial", "ConsentimientoDigital"])
    for i in range(1, 4):
        med.append([f"M{i:03d}", f"Médico {i}", "Clínica", "Ciudad X", "A", "Alto", "Sí"])
    prod = wb.create_sheet("Producto")
    prod.append(["ID", "Producto", "Molecula", "ATC4", "AreaTerapeutica", "Corporacion"])
    for i in range(1, 4):
        prod.append([f"P{i:03d}", f"Producto {i}", f"Mol {i}", "A01A", "Área", "Corp Z"])
    fec = wb.create_sheet("Fecha")
    fec.append(["Fecha", "Anio", "MesNumero", "Mes", "AnioMes", "Trimestre", "DiaSemanaNumero"])
    fec.append(["2025-01-01", 2025, 1, "Ene", "2025-01", 1, 3])
    vis = wb.create_sheet("Visitas")
    vis.append(["ID", "Fecha", "MedicoID", "ProductoID", "APM", "Estado", "Visitas"])
    vis.append([1, "2025-01-02", "M001", "P001", "APM 1", "Hecha", 1])
    rec = wb.create_sheet("Recetas")
    rec.append(["ID", "Fecha", "MedicoID", "ProductoID", "PXs", "Especialidad"])
    rec.append([1, "2025-01-03", "M001", "P001", 4, "Clínica"])
    dig = wb.create_sheet("Interacc. Digital")
    dig.append(["InteraccionID", "Fecha", "MedicoID", "ProductoID", "Canal", "Enviado",
                "Abierto", "Click", "DuracionMin", "Especialidad"])
    dig.append([1, "2025-01-04", "M001", "P002", "Mail", 1, 1, 0, 2.5, "Clínica"])
    if con_mercado:
        mer = wb.create_sheet("Mercado")
        mer.append(["Fecha", "ProductoID", "Corporacion", "Unidades", "VentasUSD"])
        mer.append(["2025-01-01", "P001", "Corp Z", 10, 100.0])
        mer.append(["2025-02-01", "P001", "Corp Z", 30, 500.0])   # P001: 600/40 = 15
        mer.append(["2025-01-01", "P002", "Corp Z", 4, 50.0])     # P002: 12.5
        mer.append(["2025-02-01", "P003", "Corp Z", 0, 20.0])     # sin unidades: sin precio
    wb.save(ruta)


def test_ventas_del_libro_de_ocho_hojas_salen_de_mercado(tmp_path):
    ruta = tmp_path / "bases.xlsx"
    _libro_ocho_hojas(ruta)
    ev = conectores.evaluar_hojas(str(ruta), "ventas")
    assert conectores.mejor_hoja(ev) == "Mercado"
    v = conectores.leer_archivo(str(ruta), "ventas")
    assert v.attrs["hoja"] == "Mercado"
    m = conectores.autodetectar_mapeo(v, "ventas")
    assert m == {"Fecha": "fecha", "ProductoID": "sku", "Unidades": "cantidad"}
    out = conectores.normalizar(v, "ventas", m)
    assert len(out) == 4 and out["cantidad"].sum() == 44
    # Sin cliente ni comprobante en la hoja: las columnas existen VACÍAS,
    # para que la analítica no tire KeyError, pero no se inventa ningún id.
    assert out["cliente_id"].isna().all() and out["venta_id"].isna().all()


def test_las_recetas_no_cuentan_como_ventas(tmp_path):
    # «PXs» no es sinónimo de cantidad: con él, «Recetas» empataría con
    # «Mercado» y ganaría por venir antes, dejando ventas sin monto.
    ruta = tmp_path / "bases.xlsx"
    _libro_ocho_hojas(ruta)
    ev = {e["hoja"]: e for e in conectores.evaluar_hojas(str(ruta), "ventas")}
    assert ev["Recetas"]["faltan"] == ["cantidad"]
    assert ev["Mercado"]["faltan"] == []


def test_hoja_de_ventas_mal_elegida_nombra_la_que_sirve(tmp_path):
    ruta = tmp_path / "bases.xlsx"
    _libro_ocho_hojas(ruta)
    ev = conectores.evaluar_hojas(str(ruta), "ventas")
    assert conectores.hoja_completa_alternativa(ev, "Producto") == "Mercado"
    assert conectores.hoja_completa_alternativa(ev, "Recetas") == "Mercado"
    assert conectores.hoja_completa_alternativa(ev, "Mercado") is None
    # Si ninguna hoja sirve entera, no hay alternativa que ofrecer.
    sin = tmp_path / "sin_mercado.xlsx"
    _libro_ocho_hojas(sin, con_mercado=False)
    ev_sin = conectores.evaluar_hojas(str(sin), "ventas")
    assert conectores.hoja_completa_alternativa(ev_sin, "Producto") is None


def test_productos_toman_el_precio_de_la_hoja_mercado_del_libro(tmp_path):
    ruta = tmp_path / "bases.xlsx"
    _libro_ocho_hojas(ruta)
    p = conectores.leer_archivo(str(ruta), "productos")
    assert p.attrs["hoja"] == "Producto"
    m = conectores.autodetectar_mapeo(p, "productos")
    assert m == {"ID": "sku", "Producto": "nombre"}
    assert conectores.faltan_obligatorias(p, "productos", m) == ["precio"]
    # Aunque la hoja elegida para ventas fuera otra, el libro trae Mercado.
    precio, nota, hoja = conectores.precio_desde_libro(str(ruta))
    assert hoja == "Mercado"
    assert precio.to_dict() == {"P001": 15.0, "P002": 12.5}   # P003: 0 unidades
    assert "DERIVADO" in nota and "Mercado" in nota
    out = conectores.normalizar(conectores.completar_precio(p, m, precio), "productos", m)
    por_sku = out.set_index("sku")
    assert por_sku.loc["P001", "nombre"] == "Producto 1"
    assert por_sku.loc["P001", "precio"] == 15.0
    assert por_sku.loc["P002", "precio"] == 12.5
    assert por_sku.loc["P003", "precio"] == 0.0   # sin dato: no se inventa


def test_sin_hoja_con_monto_no_hay_precio_del_libro(tmp_path):
    ruta = tmp_path / "sin_mercado.xlsx"
    _libro_ocho_hojas(ruta, con_mercado=False)
    assert conectores.precio_desde_libro(str(ruta)) == (None, "", None)


def test_el_codigo_cruza_aunque_una_hoja_lo_traiga_como_decimal():
    ventas = pd.DataFrame({"fecha": ["2025-01-01"] * 2, "sku": [101.0, 102.0],
                           "cantidad": [2, 4], "VentasUSD": [20.0, 10.0]})
    precio, _ = conectores.precio_desde_ventas(ventas)
    productos = pd.DataFrame({"sku": [101, 102], "nombre": ["a", "b"]})
    out = conectores.completar_precio(productos, {}, precio)
    assert out["precio"].tolist() == [10.0, 2.5]


def test_ventas_sin_cliente_no_rompen_la_analitica(tmp_path):
    from plania import analitica, sugerencias
    ruta = tmp_path / "bases.xlsx"
    _libro_ocho_hojas(ruta)
    v = conectores.leer_archivo(str(ruta), "ventas")
    ventas = conectores.normalizar(v, "ventas")
    p = conectores.leer_archivo(str(ruta), "productos")
    mp = conectores.autodetectar_mapeo(p, "productos")
    precio, _, _ = conectores.precio_desde_libro(str(ruta))
    productos = conectores.normalizar(conectores.completar_precio(p, mp, precio),
                                      "productos", mp)
    clientes = conectores.normalizar(conectores.leer_archivo(str(ruta), "clientes"),
                                     "clientes")
    datos = {"productos": productos, "clientes": clientes, "ventas": ventas}
    enr = analitica.enriquecer_ventas(ventas, productos, clientes)
    k = analitica.kpis(productos, enr)
    assert k["clientes_activos"] == 0          # no hay dato, no un número falso
    sugerencias.generar_todas(datos)           # no levanta KeyError('cliente_id')
    # Y sobrevive al ida y vuelta por la base SQLite que usa la app.
    url = conectores.guardar_como_base(datos, str(tmp_path / "erp.db"))
    vuelta = conectores.cargar_datos(url)
    assert len(vuelta["ventas"]) == 4 and "cliente_id" in vuelta["ventas"].columns


def test_la_pantalla_ofrece_la_hoja_buena_y_el_precio_del_libro():
    fuente = open(APP, encoding="utf-8").read()
    assert "conectores.hoja_completa_alternativa(" in fuente
    assert "on_click=_elegir_hoja" in fuente
    assert "conectores.precio_desde_libro(" in fuente
